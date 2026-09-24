"""Does the ported Excel report agree with One Supply, and open as a workbook?

The value formatting is a port of Bubble's BaseExporter._format_special_value,
so the test is not "does it look right" but "does it agree with the original".
This runs the real Python side by side with the C# port's output, then opens
the downloaded workbook and checks the things a spreadsheet silently gets wrong:
tolerances coerced to numbers, and GD&T glyphs left in a face that has no glyph
for them.

    python .mssql-scripts/check_balloon_excel.py [part_id]
"""

import io
import os
import re
import sys

BUBBLE = r"C:\Aizera\RPA\Bubble\(公共版)pyqt-bubble--api\(公共版)pyqt-bubble--api"
PART = sys.argv[1] if len(sys.argv) > 1 else "12"
BASE = "http://localhost:5001"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "shots")
os.makedirs(OUT, exist_ok=True)

import openpyxl                                                 # noqa: E402
from playwright.sync_api import sync_playwright                 # noqa: E402

fails = []


def check(name, ok, detail=""):
    print(f"  {'PASS' if ok else 'FAIL'}  {name}{'  ' + str(detail) if detail else ''}")
    if not ok:
        fails.append(name)


# ── the original, lifted out of BaseExporter without importing PySide6 ──────
src = io.open(os.path.join(BUBBLE, "core", "exporters", "base_exporter.py"),
              encoding="utf-8").read()
ns = {"re": re}
for name in ("_fix_leading_zero_num", "_format_special_value"):
    body = re.search(rf"\n    (?:@staticmethod\n    )?def {name}\(.*?"
                     r"(?=\n    (?:@|def |class ))", src, re.S).group(0)
    exec(compile("\n".join(l[4:] for l in body.split("\n")), name, "exec"), ns)  # noqa: S102


def original(value):
    """Bubble's answer, with smart formatting on - what the port must match."""
    class _Self:
        _fix_leading_zero_num = staticmethod(ns["_fix_leading_zero_num"])
    return ns["_format_special_value"](_Self(), value, enable_smart_format=True)


print("1. the port agrees with One Supply on real drawing values")
CASES = [".380", "-.002", "+.005", "0.653", "4X .100/.097", ".100/.097",
         "1/2", "0.015-0.010", "2X \u2300.380", "1.250", "", "45\u00b0\u00b11\u00b0",
         "\u2316", ".938"]
for v in CASES:
    want = original(v)
    print(f"        {v!r:<22} -> {want!r}")
# The C# side is exercised through the real endpoint below; here we only pin
# the expectations, so a change to either side shows up as a diff.
expected = {v: original(v) for v in CASES}
check("leading zero added", expected[".380"] == "0.380")
check("negative leading zero", expected["-.002"] == "-0.002")
check("limit pair reordered small-first",
      expected["4X .100/.097"] == "4X 0.097-0.100", expected["4X .100/.097"])
check("a fraction is left alone", expected["1/2"] == "1/2")
check("a whole number is untouched", expected["1.250"] == "1.250")

print("\n2. the toolbar button downloads a workbook")
# Through the browser rather than requests: the endpoint is cookie-authorised
# and, more to the point, the button is half of what is being tested. A session
# forged with requests would prove the endpoint works and say nothing about
# whether anything reaches it.
path = os.path.join(OUT, f"balloons-{PART}.xlsx")

with sync_playwright() as p:
    browser = p.chromium.launch(downloads_path=OUT)
    page = browser.new_page(viewport={"width": 1600, "height": 950},
                            accept_downloads=True)
    page.goto(f"{BASE}/Account/Login", wait_until="networkidle")
    page.wait_for_timeout(1500)
    user = page.get_by_placeholder("user name")
    if user.count():
        user.fill("admin")
        page.get_by_placeholder("password").fill("serenity")
        page.get_by_role("button", name="Sign In").click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(2000)

    page.goto(f"{BASE}/Costing/Workspace/{PART}", wait_until="networkidle")
    page.wait_for_timeout(3000)
    doc2d = page.locator(".cw-rail-docs [class*=row]:has-text('2D')").first
    if not doc2d.count():
        doc2d = page.locator(".cw-rail-docs >> text=2D").first
    if doc2d.count():
        doc2d.click()
        page.wait_for_timeout(3500)
    mode = page.locator(".cw-modes button:has-text('Balloon')").first
    if mode.count():
        mode.click()
    page.wait_for_timeout(6000)

    btn = page.locator("#btn-export-xls")
    check("the Export Excel button is in the toolbar", btn.count() > 0)
    check("it is enabled", btn.count() > 0 and btn.is_enabled())

    if btn.count():
        with page.expect_download(timeout=120000) as dl:
            btn.click()
        download = dl.value
        download.save_as(path)
        print(f"        suggested name: {download.suggested_filename}")
        check("the download is named for the part",
              download.suggested_filename.startswith("Balloons_")
              and download.suggested_filename.endswith(".xlsx"),
              download.suggested_filename)
    browser.close()

if not os.path.exists(path):
    print(f"\n{len(fails)} FAILED")
    sys.exit(1)
print(f"        wrote {path} ({os.path.getsize(path)} bytes)")

print("\n3. what is actually in it")
wb = openpyxl.load_workbook(path)
ws = wb.active
headers = [c.value for c in ws[2]]
print(f"        headers: {headers}")
check("the tolerance columns are there",
      "Upper Tol" in headers and "Lower Tol" in headers, headers)
check("the characteristic column is there", "Characteristic" in headers)
check("the GD&T glyph has its own column", "GD&T" in headers, headers)
rows = list(ws.iter_rows(min_row=3, values_only=True))
print(f"        {len(rows)} data rows")
check("it has rows", len(rows) > 0)

# The failure that matters: Excel reading "0.380" as a float and dropping the
# trailing zero, or "1-2" as a date. Both destroy a tolerance silently.
typed = [r for r in rows if any(isinstance(v, (int, float)) for v in r[4:6])]
check("no tolerance was coerced to a number", not typed, typed[:2])

tol_rows = [r for r in rows if (r[4] or "") or (r[5] or "")]
print(f"        {len(tol_rows)} rows carry a tolerance")
for r_ in tol_rows[:4]:
    print(f"          {r_[0]!r:<6} gdt={r_[1]!r:<4} sym={r_[3]!r:<16} up={r_[4]!r:<9} lo={r_[5]!r}")
check("tolerances made it into the sheet", len(tol_rows) > 0)
# And they must be formatted the way One Supply formats them.
bad = [r_ for r_ in tol_rows
       if r_[4] and original(str(r_[4])) != str(r_[4])
       or r_[5] and original(str(r_[5])) != str(r_[5])]
check("every tolerance matches One Supply's formatting", not bad, bad[:2])

# Glyphs: a pure rare symbol must have been re-faced, or it renders as a box.
RARE = set("\u2316\u232f\u23e4\u2334\u25b1\u232d\u2312\u2313\u2330\u2300")
pure = [(c.row, c.value, c.font.name) for row in ws.iter_rows(min_row=3)
        for c in row if c.value and str(c.value).strip() in RARE]
print(f"        {len(pure)} pure-glyph cells")
check("there ARE pure-glyph cells to style", len(pure) > 0, len(pure))
check("pure glyph cells use a face that has the glyph",
      all(f == "Segoe UI Symbol" for _r, _v, f in pure), pure[:3])
check("and were bumped to a legible size",
      all((ws.cell(row=r_, column=2).font.size or 0) >= 14 for r_, _v, _f in pure),
      [(r_, ws.cell(row=r_, column=2).font.size) for r_, _v, _f in pure[:3]])
# ...and a mixed cell must NOT have been bolded up to 14pt, which would read as
# a defect beside its neighbours.
mixed = [c for row in ws.iter_rows(min_row=3) for c in row
         if c.value and len(str(c.value).strip()) > 2
         and any(ch in RARE for ch in str(c.value))]
check("mixed text was not blown up",
      all((c.font.size or 11) <= 12 for c in mixed),
      [(c.value, c.font.size) for c in mixed[:2]])

print(f"\n{'ALL PASS' if not fails else str(len(fails)) + ' FAILED: ' + ', '.join(fails)}")
sys.exit(1 if fails else 0)
